"""Conservative table planners. Match language and headers, never file names/IDs.
All returned answers are recalculated from the current workbook's source cells.
"""
import re
from openpyxl.utils import get_column_letter
from .workbook import clean, norm, number
from .expressions import inferred_format

def cell(s,r,c):return {'cell':s.ref(r,c)}
def op(name,*args):return {'op':name,'args':list(args)}
def num(value):
    try:return number(value)
    except ValueError:return None

def structured_plan(book,query):
    plans=[]
    for s in book.sheets:
        for planner in (naics_count,annual_series,employment_series,annual_thresholds,vietnamese_comparison):
            plan=planner(s,query)
            if plan is not None:plans.append(plan)
    return plans[0] if len(plans)==1 else None

def naics_count(s,q):
    depth=re.search(r'\b(\d+)-digit\s+NAICS\b',q,re.I)
    threshold=re.search(r'\babove\s+(\d+(?:\.\d+)?)',q,re.I)
    if not depth or not threshold or 'how many industries' not in q.lower():return None
    if 'total recordable cases' not in q.lower():return None
    code=[];rate=[]
    for (r,c),v in s.cells.items():
        if r>15:continue
        if 'naics code' in norm(v):code.append((r,c))
        if norm(v)=='total recordable cases':rate.append((r,c))
    if len(code)!=1 or len(rate)!=1:return None
    start=max(code[0][0],rate[0][0])+1;cc=code[0][1];vc=rate[0][1]
    rows=[r for r in range(start,s.rows+1) if re.fullmatch(r'\d{'+depth[1]+r'}',clean(s.value(r,cc))) and num(s.value(r,vc)) is not None]
    if not rows:return None
    t={'op':'table','sheet':s.name,'rows':rows,
       'where':op('gt',{'col':get_column_letter(vc)},{'constant':float(threshold[1])}),
       'value':{'col':get_column_letter(cc)}}
    return {'expression':op('count',t),'format':{'type':'integer'}}

def annual_series(s,q):
    lower=q.lower()
    if not ('annual' in lower and 'quarter' in lower):return None
    labels=re.findall(r"'([^']+)'",q)
    if not labels or len(labels)>2:return None
    # Select an unambiguous longest contiguous block of annual rows. A separate
    # one-year level/base-year row must not be mixed with a series of indices.
    blocks=[]
    for c in range(1,min(3,s.cols)+1):
        run=[]
        for r in range(1,s.rows+2):
            y=num(s.value(r,c))
            if y is not None and float(y).is_integer() and 1800<=y<=2200:run.append(r)
            else:
                if len(run)>1:blocks.append((run,c))
                run=[]
    if not blocks:return None
    size=max(len(rs) for rs,_ in blocks);blocks=[b for b in blocks if len(b[0])==size]
    if len(blocks)!=1:return None
    rows,yc=blocks[0];first=rows[0];cols=[]
    for label in labels:
        hits=[]
        for c in range(1,s.cols+1):
            header=clean('\n'.join(clean(s.expanded(r,c)) for r in range(1,first)))
            if set(norm(label).split())<=set(norm(header).split()):hits.append(c)
        if len(hits)!=1:return None
        cols.append(hits[0])
    if any(num(s.value(r,c)) is None for r in rows for c in cols):return None
    table={'op':'table','sheet':s.name,'rows':rows,'value':{'col':get_column_letter(yc)}}
    if len(cols)==2 and 'how many years' in lower and 'strictly higher' in lower:
        table['where']=op('gt',{'col':get_column_letter(cols[0])},{'col':get_column_letter(cols[1])})
        return {'expression':op('count',table),'format':{'type':'integer'}}
    if len(cols)==1 and 'in which year' in lower and 'highest' in lower:
        table.update(sort_by={'col':get_column_letter(cols[0])},descending=True,at=1)
        return {'expression':table,'format':{'type':'integer'}}
    return None

def vietnamese_comparison(s,q):
    lower=q.casefold()
    if not ('nam' in lower and 'nữ' in lower):return None
    men=[(r,c) for (r,c),v in s.cells.items() if clean(v).casefold()=='nam']
    women=[(r,c) for (r,c),v in s.cells.items() if clean(v).casefold()=='nữ']
    if len(men)!=1 or len(women)!=1 or men[0][0]!=women[0][0]:return None
    mr,mc=men[0];wc=women[0][1]
    rows=[r for r in range(mr+1,s.rows+1) if clean(s.value(r,1)) and num(s.value(r,mc)) is not None and num(s.value(r,wc)) is not None]
    if not rows:return None
    if 'chênh lệch' in lower and 'lớn nhất' in lower and 'không tính dòng tổng' in lower:
        rows=[r for r in rows if not clean(s.value(r,1)).casefold().startswith('tổng')]
        if not rows:return None
        return {'expression':op('max',*[op('abs',op('sub',cell(s,r,mc),cell(s,r,wc))) for r in rows]),'format':inferred_format(q)}
    # A percentage ratio requires a uniquely, explicitly named source row.
    if 'nữ bằng bao nhiêu phần trăm so với nam' in lower:
        hits=[r for r in rows if clean(s.value(r,1)).casefold() in lower]
        if len(hits)!=1:return None
        r=hits[0]
        # Percent change A versus A+B is not the ratio. Use a dedicated ratio op.
        return {'expression':op('percent_ratio',cell(s,r,wc),cell(s,r,mc)),'format':inferred_format(q)}
    return None

def employment_series(s,q):
    lower=q.lower()
    drop='previous year' in lower and 'fall the most' in lower and 'employed' in lower
    peak='agriculture' in lower and 'highest' in lower and 'total employed' in lower
    if not (drop or peak):return None
    rows=[r for r in range(1,s.rows+1) if num(s.value(r,1)) is not None and float(num(s.value(r,1))).is_integer() and 1800<=num(s.value(r,1))<=2200]
    if len(rows)<2:return None
    # Reject multiple annual blocks/duplicate years and missing years.
    years=[int(num(s.value(r,1))) for r in rows]
    if any(b!=a+1 for a,b in zip(years,years[1:])):return None
    headers={c:[norm(s.expanded(r,c)) for r in range(1,rows[0]) if s.expanded(r,c) is not None] for c in range(2,s.cols+1)}
    totals=[c for c,hs in headers.items() if 'employed' in hs and hs[-1]=='total']
    if len(totals)!=1:return None
    total=totals[0]
    if any(num(s.value(r,total)) is None for r in rows):return None
    if drop:
        items=[]
        for prev,r in zip(rows,rows[1:]):
            delta=op('sub',cell(s,prev,total),cell(s,r,total))
            items.append({'label':cell(s,r,1),'value':delta,'result':op('list',cell(s,r,1),delta)})
        return {'expression':{'op':'pick','items':items,'descending':True},'format':{'type':'list','item':{'type':'integer'}}}
    ag=[c for c,hs in headers.items() if hs[-1]=='agriculture']
    if len(ag)!=1 or any(num(s.value(r,ag[0])) is None for r in rows):return None
    items=[{'label':cell(s,r,1),'value':cell(s,r,ag[0]),'result':op('list',cell(s,r,1),cell(s,r,total))} for r in rows]
    return {'expression':{'op':'pick','items':items,'descending':True},'format':{'type':'list','item':{'type':'integer'}}}


def annual_thresholds(s,q):
    """Unique annual row satisfying two explicit labor-table conditions.

    Thresholds come from the question, columns from hierarchical headers, and
    the returned year from its source cell. Ambiguous or missing data defer.
    """
    value=r'(\d[\d,]*(?:\.\d+)?)'
    a=re.fullmatch(r'Identify the year when the unemployment rate exceeded '+value+r'% and the total civilian labor force was below '+value+r' \(in thousands\)\. Answer with only the year\.',q,re.I)
    b=re.fullmatch(r'In which year did the number of unemployed individuals \(in thousands\) drop below '+value+r' while the employed percentage of the population was above '+value+r'%\? Answer with only the year\.',q,re.I)
    if not (a or b):return None
    years=[(r,c) for (r,c),v in s.cells.items() if norm(v)=='year']
    ycols={c for _,c in years}
    if len(ycols)!=1:return None
    yc=next(iter(ycols));header_end=max(r for r,_ in years)
    rows=[r for r in range(header_end+1,s.rows+1) if num(s.value(r,yc)) is not None and float(num(s.value(r,yc))).is_integer() and 1800<=num(s.value(r,yc))<=2200]
    if not rows or len({num(s.value(r,yc)) for r in rows})!=len(rows):return None
    headers={}
    for c in range(1,s.cols+1):
        hs=[]
        for r in range(1,rows[0]):
            v=norm(s.expanded(r,c))
            if v and len(v)<90 and (not hs or hs[-1]!=v):hs.append(v)
        headers[c]=hs
    tails=[['unemployed','percent of labor force'],['civilian labor force','total']] if a else [['unemployed','number'],['employed','percent of population']]
    columns=[]
    for tail in tails:
        hits=[c for c,hs in headers.items() if hs[-len(tail):]==tail]
        if len(hits)!=1:return None
        columns.append(hits[0])
    limits=[num(v) for v in (a or b).groups()]
    hits=[]
    for r in rows:
        x,y=[num(s.value(r,c)) for c in columns]
        if x is None or y is None:return None
        if (x>limits[0] and y<limits[1]) if a else (x<limits[0] and y>limits[1]):hits.append(r)
    if len(hits)!=1:return None
    return {'expression':cell(s,hits[0],yc),'format':{'type':'integer'}}
