import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from excelqa.workbook import Book
from excelqa.local import local_plan
from excelqa.expressions import Evaluator, format_answer

class LocalTests(unittest.TestCase):
    def solve(self, rows, q, offset=0, merges=()):
        with tempfile.TemporaryDirectory() as d:
            w=Workbook();s=w.active;s.title='Fresh'
            for _ in range(offset):s.append([None])
            for row in rows:s.append(row)
            for rg in merges:s.merge_cells(rg)
            p=Path(d)/'new.xlsx';w.save(p);b=Book(p)
            plan=local_plan(b,q)
            self.assertIsNotNone(plan)
            return format_answer(Evaluator(b).run(plan['expression']),plan['format'])
    def test_combined_industries_uses_headers_and_changed_values(self):
        rows=[['Age','Construction','Manufacturing','Other'],['TOTAL'],['20 to 24 years',7,12,100],['16 to 19 years',1,2,50]]
        q='What is the total number of employed persons aged 20 to 24 years in Construction and Manufacturing combined? Answer with only the integer.'
        self.assertEqual(self.solve(rows,q), '19')
        self.assertEqual(self.solve(rows,q,offset=4), '19')
    def test_group_selection_does_not_use_men_for_women(self):
        rows=[['Status','School A','School B'],['Men'],['Unemployment rate',8,10],['Women'],['Unemployment rate',2,4]]
        q='What was the average unemployment rate across all education levels for women? Answer with only the number rounded to two decimal places.'
        self.assertEqual(self.solve(rows,q),'3.00')
    def test_unknown_question_declines_instead_of_guessing(self):
        with tempfile.TemporaryDirectory() as d:
            w=Workbook();w.active.append(['Mystery',1]);p=Path(d)/'a.xlsx';w.save(p)
            self.assertIsNone(local_plan(Book(p),'Explain what will happen next year'))

if __name__=='__main__':unittest.main()

class RegressionTests(unittest.TestCase):
    solve = LocalTests.solve
    def test_nested_subtraction_keeps_decimal_type(self):
        rows=[['Age','Construction','Manufacturing'],['TOTAL'],['20 to 24 years',7,12],['16 to 19 years',1,2]]
        self.assertEqual(self.solve(rows,'How many more employed people were aged 20 to 24 than aged 16 to 19 across all industries? Answer with only the integer.'),'16')
    def test_grouped_column_does_not_add_nonagriculture(self):
        rows=[['Age','Agriculture and related industries','Nonagricultural industries'],[None,'Unpaid family workers','Unpaid family workers'],['Total, 16 years and over',35,50]]
        self.assertEqual(self.solve(rows,'Calculate the total number of unpaid family workers across all age groups in the agricultural sector. What is the final sum? Answer with only the integer.'),'35')
    def test_total_column_not_added_to_subcategories(self):
        rows=[['Industry','Total employed','Protective service occupations','Service occupations, except protective'],['TOTAL'],['Manufacturing',1000,45,220]]
        self.assertEqual(self.solve(rows,"Calculate the combined total of employees in 'Protective service occupations' and 'Service occupations, except protective' for the Manufacturing sector. Answer with only the integer."),'265')
    def test_multi_values_use_parent_column_labels(self):
        rows=[['Characteristic','Total at work','Worked 1 to 34 hours','Worked 35 hours or more'],[None,None,'Total',None],['Total, 16 years and over',100,30,70]]
        self.assertEqual(self.solve(rows,'What are number of people who worked 1 to 34 hours and worked 35 hours or more. Answer as two integers, separated by a comma and one space.'),'30, 70')

    def test_parent_match_does_not_select_unrequested_subcolumns(self):
        rows=[['Characteristic','Worked 1 to 34 hours',None,'Worked 35 hours or more'],[None,'Total','For economic reasons',None],['Total, 16 years and over',30,5,70]]
        self.assertEqual(self.solve(rows,'What are number of people who worked 1 to 34 hours and worked 35 hours or more. Answer as two integers, separated by a comma and one space.',merges=('B1:C1',)),'30, 70')

    def test_column_comparison_preserves_question_order(self):
        rows=[['Age','Transportation and utilities','Financial activities'],['TOTAL'],['20 years and over',10,20]]
        self.assertEqual(self.solve(rows,'For persons aged 20 years and over, does employment in Financial Activities exceed that in Transportation and Utilities? Answer with only Yes or No.'),'Yes')

    def test_agriculture_total_not_added_to_unpaid_workers(self):
        rows=[['Age','Agriculture and related industries',None,'Nonagricultural industries',None],
              [None,'Total','Unpaid family workers','Total','Unpaid family workers'],
              ['Total, 16 years and over',2264,35,10000,50]]
        self.assertEqual(self.solve(rows,'Calculate the total number of unpaid family workers across all age groups in the agricultural sector. What is the final sum? Answer with only the integer.',merges=('B1:C1','D1:E1')),'35')

    def test_numeric_row_index_not_returned_as_category_name(self):
        with tempfile.TemporaryDirectory() as d:
            w=Workbook();s=w.active;s.append(['Indent Level','Category','Percent change']);s.append([0,'A',1]);s.append([1,'B',6])
            p=Path(d)/'x.xlsx';w.save(p)
            self.assertIsNone(local_plan(Book(p),'Which category has highest percent change?'))
